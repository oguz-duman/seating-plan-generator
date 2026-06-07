from __future__ import annotations
import random
import shutil
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, List, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class Config:
    workbook_file: str = "config.xlsx"

    output_numbers_file: str = "seating_plan_numbers.xlsx"
    output_names_numbers_file: str = "seating_plan_names_numbers.xlsx"

    control_sheet: str = "config"
    student_number_column: str = "A"
    student_name_column: str = "B"
    class_column: str = "F"
    start_row: int = 2

    random_seed: int | None = None


@dataclass
class Student:
    number: str
    name: str


def get_column_values(ws: Worksheet, column_letter: str, start_row: int) -> List[str]:
    """Return non-empty trimmed cell values from a column starting at start_row."""

    values: List[str] = []

    for row in range(start_row, ws.max_row + 1):
        value = ws[f"{column_letter}{row}"].value

        if value is None:
            continue

        text = str(value).strip()
        if text:
            values.append(text)

    return values


def get_students(
    ws: Worksheet,
    number_column: str,
    name_column: str,
    start_row: int,
) -> List[Student]:
    """Parse students from worksheet rows, validating completeness of number/name."""

    students: List[Student] = []

    for row in range(start_row, ws.max_row + 1):
        number_value = ws[f"{number_column}{row}"].value
        name_value = ws[f"{name_column}{row}"].value

        number_text = "" if number_value is None else str(number_value).strip()
        name_text = "" if name_value is None else str(name_value).strip()

        if "(" in name_text:
            # remove extra info like "(useR_name)" in student names
            name_text = name_text.split("(")[0].strip()  

        if number_text == "" and name_text == "":
            continue

        if number_text == "" or name_text == "":
            raise ValueError(
                f"Incomplete student data at row {row}: "
                f"{number_column}{row}='{number_text}', {name_column}{row}='{name_text}'"
            )

        students.append(Student(number=number_text, name=name_text))

    return students


def find_seat_cells(ws: Worksheet) -> List[Tuple[int, int]]:
    """Return coordinates of cells containing 1, representing available seats."""
    
    seats: List[Tuple[int, int]] = []

    for row in ws.iter_rows():
        for cell in row:
            # cell value 1 marks available seat in layout template
            if cell.value == 1 or str(cell.value).strip() == "1":
                assert cell.row is not None
                assert cell.column is not None
                seats.append((cell.row, cell.column))
    return seats


def assign_students_round_robin(
    students: List[Student],
    class_seats: dict[str, List[Tuple[int, int]]],
    class_order: List[str],
) -> dict[str, List[Tuple[Tuple[int, int], Student]]]:
    """Distribute students across classes in round-robin order based on seat availability."""

    assignments: dict[str, List[Tuple[Tuple[int, int], Student]]] = {
        class_name: [] for class_name in class_order
    }

    seat_indices = {class_name: 0 for class_name in class_order}
    active_classes = [c for c in class_order if len(class_seats[c]) > 0]

    if not active_classes:
        return assignments

    class_pointer = 0

    for student in students:
        if not active_classes:
            break

        assigned = False
        checked_count = 0

        while checked_count < len(active_classes):
            class_name = active_classes[class_pointer]
            seat_list = class_seats[class_name]
            seat_idx = seat_indices[class_name]

            if seat_idx < len(seat_list):
                seat_coord = seat_list[seat_idx]
                assignments[class_name].append((seat_coord, student))
                seat_indices[class_name] += 1
                assigned = True

                class_pointer = (class_pointer + 1) % len(active_classes)       # cycle through classes evenly
                break

            checked_count += 1
            class_pointer = (class_pointer + 1) % len(active_classes)           

        active_classes = [
            c for c in active_classes if seat_indices[c] < len(class_seats[c])
        ]

        if active_classes:
            class_pointer %= len(active_classes)

        if not assigned and not active_classes:
            break

    return assignments


def prepare_workbook(source_path: Path, class_names: List[str], exam_name: str, control_sheet_name: str):
    """Load workbook, keep only relevant class sheets, and detect seat coordinates."""

    wb = load_workbook(source_path)

    if control_sheet_name not in wb.sheetnames:
        raise ValueError(f"Control sheet not found in workbook: {control_sheet_name}")

    missing_sheets = [name for name in class_names if name not in wb.sheetnames]
    if missing_sheets:
        raise ValueError(f"These class sheets do not exist in workbook: {missing_sheets}")

    for class_name in class_names:
        ws = wb[class_name]
        ws["A1"] = f"{exam_name} | {class_name}"

    used_sheet_names = set(class_names)

    for sheet_name in wb.sheetnames[:]:
        if sheet_name not in used_sheet_names:
            del wb[sheet_name]                      # keep only classroom layout sheets

    class_seats: dict[str, List[Tuple[int, int]]] = {}
    total_capacity = 0

    for class_name in class_names:
        ws = wb[class_name]
        seats = find_seat_cells(ws)
        class_seats[class_name] = seats
        total_capacity += len(seats)

    return wb, class_seats, total_capacity


def clear_remaining_ones(wb, class_names: List[str]) -> None:
    """Remove unused seat markers (cells still containing 1)."""

    for class_name in class_names:
        ws = wb[class_name]

        for row in ws.iter_rows():
            for cell in row:
                if cell.value == 1 or cell.value == "1":
                    cell.value = None


def fill_workbook(
    wb,
    assignments: dict[str, List[Tuple[Tuple[int, int], Student]]],
    class_names: List[str],
    formatter: Callable[[Student], str],
) -> int:
    """Fill seat cells using formatter and enable text wrapping."""

    placed_students = 0

    for class_name, seat_assignments in assignments.items():
        ws = wb[class_name]
        for (row, col), student in seat_assignments:
            cell = ws.cell(row=row, column=col)
            cell.value = formatter(student)

            new_alignment = copy(cell.alignment)
            new_alignment.wrap_text = True              # allow number + name to appear on separate lines
            cell.alignment = new_alignment

            placed_students += 1

    clear_remaining_ones(wb, class_names)
    return placed_students


def sanitize_folder_name(name: str) -> str:
    """Replace filesystem-invalid characters for safe folder creation."""

    invalid_chars = '<>:"/\\|?*'
    sanitized = "".join("_" if ch in invalid_chars else ch for ch in name).strip()
    sanitized = sanitized.rstrip(". ")
    if not sanitized:
        raise ValueError("Exam name is empty or invalid for folder creation.")
    return sanitized


def prepare_output_folder(folder_path: Path) -> None:
    """Create or clean output folder before writing new files."""

    if folder_path.exists():
        if not folder_path.is_dir():
            raise RuntimeError(f"Output path exists but is not a folder: {folder_path}")

        for item in folder_path.iterdir():
            try:
                if item.is_dir():
                    shutil.rmtree(item)
                else:
                    item.unlink()
            except Exception as e:
                raise RuntimeError(
                    f"Could not delete existing item in output folder: {item}. "
                    "Make sure no file inside the folder is open."
                ) from e
    else:
        try:
            folder_path.mkdir(parents=True, exist_ok=False)
        except Exception as e:
            raise RuntimeError(f"Could not create output folder: {folder_path}") from e


def main() -> None:
    """Generate randomized seating plans from config.xlsx and save formatted outputs."""

    config = Config(
        workbook_file="config.xlsx",
        output_numbers_file="seating_plan_numbers.xlsx",
        output_names_numbers_file="seating_plan_names_numbers.xlsx",
        control_sheet="config",
        student_number_column="A",
        student_name_column="B",
        class_column="F",
        start_row=2,
        random_seed=42,
    )

    if config.random_seed is not None:
        random.seed(config.random_seed)         # ensures reproducible seating plan

    workbook_path = Path(config.workbook_file)

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook file not found: {workbook_path}")

    control_wb = load_workbook(workbook_path)
    if config.control_sheet not in control_wb.sheetnames:
        raise ValueError(f"Control sheet not found: {config.control_sheet}")

    control_ws = control_wb[config.control_sheet]

    exam_name_value = control_ws["D2"].value
    exam_name = "" if exam_name_value is None else str(exam_name_value).strip()

    if not exam_name:
        raise ValueError("Exam name in D2 is empty.")

    students = get_students(
        control_ws,
        config.student_number_column,
        config.student_name_column,
        config.start_row,
    )
    class_names = get_column_values(control_ws, config.class_column, config.start_row)

    if not students:
        raise ValueError("No students found in control sheet.")

    if not class_names:
        raise ValueError("No class names found in control sheet.")

    random.shuffle(students)

    workbook_for_numbers, class_seats, total_capacity = prepare_workbook(
        workbook_path,
        class_names,
        exam_name,
        config.control_sheet,
    )

    if total_capacity == 0:
        raise ValueError("No seat cells with value 1 were found.")

    assignments = assign_students_round_robin(students, class_seats, class_names)

    placed_students = sum(len(seat_assignments) for seat_assignments in assignments.values())
    unplaced_students = len(students) - placed_students

    workbook_for_names_numbers, _, _ = prepare_workbook(
        workbook_path,
        class_names,
        exam_name,
        config.control_sheet,
    )

    fill_workbook(
        workbook_for_numbers,
        assignments,
        class_names,
        formatter=lambda student: student.number,
    )

    fill_workbook(
        workbook_for_names_numbers,
        assignments,
        class_names,
        formatter=lambda student: f"{student.number}\n{student.name}",
    )

    output_root_path = workbook_path.parent / "results"
    output_folder_name = sanitize_folder_name(exam_name)
    output_folder_path = output_root_path / output_folder_name
    prepare_output_folder(output_folder_path)           # ensures clean output directory

    output_numbers_path = output_folder_path / config.output_numbers_file
    output_names_numbers_path = output_folder_path / config.output_names_numbers_file

    workbook_for_numbers.save(output_numbers_path)
    workbook_for_names_numbers.save(output_names_numbers_path)

    print("Done.")
    print(f"Students found               : {len(students)}")
    print(f"Class sheets found           : {len(class_names)}")
    print(f"Total seat capacity          : {total_capacity}")
    print(f"Students placed              : {placed_students}")
    print(f"Students unplaced            : {unplaced_students}")
    print(f"Output folder                : {output_folder_path}")
    print(f"Output numbers file          : {output_numbers_path}")
    print(f"Output names+numbers file    : {output_names_numbers_path}")


if __name__ == "__main__":
    main()